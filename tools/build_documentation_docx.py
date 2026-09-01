"""Build a Word-friendly documentation file from the Markdown template."""

import re
import sys
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "BananaCare_Documentation_Template.md"
OUTPUT = ROOT / (
    sys.argv[1] if len(sys.argv) > 1 else "BananaCare_Documentation_Template.docx"
)


def plain_text(value):
    value = value.strip()
    value = re.sub(r"\*\*(.*?)\*\*", r"\1", value)
    value = re.sub(r"`(.*?)`", r"\1", value)
    return value


def shade_cell(cell, fill):
    properties = cell._tc.get_or_add_tcPr()
    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), fill)
    properties.append(shading)


def set_cell_text(cell, value, *, bold=False, size=8):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(plain_text(value))
    run.bold = bold
    run.font.name = "Arial"
    run.font.size = Pt(size)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(document, rows):
    column_count = len(rows[0])
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    if column_count == 6:
        widths = (0.65, 1.25, 1.30, 1.60, 1.60, 0.60)
    elif column_count == 2:
        widths = (1.70, 5.30)
    else:
        widths = tuple(7.0 / column_count for _ in range(column_count))

    for row_index, values in enumerate(rows):
        for column_index, value in enumerate(values):
            cell = table.cell(row_index, column_index)
            cell.width = Inches(widths[column_index])
            set_cell_text(cell, value, bold=row_index == 0, size=7 if column_count == 6 else 9)
            if row_index == 0:
                shade_cell(cell, "D9EAD3")

    document.add_paragraph()


def add_placeholder(document, text):
    table = document.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    cell = table.cell(0, 0)
    cell.width = Inches(6.8)
    set_cell_text(cell, text, bold=True, size=10)
    shade_cell(cell, "FFF2CC")
    paragraph = cell.paragraphs[0]
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    paragraph.paragraph_format.space_before = Pt(22)
    paragraph.paragraph_format.space_after = Pt(22)
    document.add_paragraph()


def configure_document(document):
    section = document.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.65)
    section.right_margin = Inches(0.65)

    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.paragraph_format.space_after = Pt(6)

    for level in range(1, 4):
        style = document.styles[f"Heading {level}"]
        style.font.name = "Arial"
        style.font.color.rgb = RGBColor(23, 76, 54)


def build_document():
    document = Document()
    configure_document(document)
    lines = SOURCE.read_text(encoding="utf-8").splitlines()
    index = 0

    while index < len(lines):
        line = lines[index].strip()

        if not line:
            index += 1
            continue

        if line.startswith("|"):
            table_lines = []
            while index < len(lines) and lines[index].strip().startswith("|"):
                table_lines.append(lines[index].strip())
                index += 1
            rows = [
                [cell.strip() for cell in row.strip("|").split("|")]
                for row in table_lines
            ]
            rows = [row for row in rows if not all(re.fullmatch(r":?-{3,}:?", cell) for cell in row)]
            add_table(document, rows)
            continue

        if line == "---":
            document.add_page_break()
        elif line.startswith("# "):
            paragraph = document.add_heading(plain_text(line[2:]), level=1)
            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        elif line.startswith("## "):
            document.add_heading(plain_text(line[3:]), level=2)
        elif line.startswith("### "):
            document.add_heading(plain_text(line[4:]), level=3)
        elif line.startswith("> [INSERT SCREENSHOT"):
            add_placeholder(document, line[2:].strip("[]"))
        elif re.match(r"^- ", line):
            document.add_paragraph(plain_text(line[2:]), style="List Bullet")
        elif re.match(r"^\d+\. ", line):
            document.add_paragraph(plain_text(re.sub(r"^\d+\. ", "", line)), style="List Number")
        elif line.startswith("> "):
            paragraph = document.add_paragraph(plain_text(line[2:]))
            paragraph.style = "Intense Quote"
        else:
            paragraph = document.add_paragraph(plain_text(line))
            if line.startswith("**") and line.endswith("**"):
                paragraph.runs[0].bold = True

        index += 1

    document.save(OUTPUT)
    print(f"Created: {OUTPUT}")


if __name__ == "__main__":
    build_document()
